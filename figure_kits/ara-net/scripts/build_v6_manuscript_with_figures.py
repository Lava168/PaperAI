#!/usr/bin/env python3
"""Build ARA-Net V6 manuscript DOCX with inline figures and rendered equations."""
from __future__ import annotations

import argparse
import hashlib
import re
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from build_v6_manuscript_docx import (  # noqa: E402
    DARK_BLUE,
    FONT,
    GRAY,
    add_inline_runs,
    add_paragraph,
    add_table,
    parse_table,
    set_paragraph_spacing,
    set_run_font,
    setup_styles,
    strip_inline_md,
)


def resolve_repo_root() -> Path:
    for cand in [Path("."), SCRIPT_DIR.parent, Path("chapter1_foundation/ARA-Net")]:
        md = cand / "reports/v6_final_model/manuscript_v6_full_draft.md"
        if md.exists():
            return cand.resolve()
    return SCRIPT_DIR.parent.resolve()


def figure_assets(root: Path) -> dict[int, list[tuple[Path, float]]]:
    v6 = root / "reports/v6_final_model/figures"
    algo = root / "reports/v6_algorithm_innovation/figures"
    brain = root / "reports/brain_figures_fcstyle/figures"
    return {
        1: [(v6 / "figure1_overall_evidence_chain_study_design.png", 6.3)],
        2: [(v6 / "figure2_atlas_guided_multimodal_feature_system.png", 6.3)],
        3: [(v6 / "figure3_rcspe_architecture_nbe_style.png", 6.3)],
        4: [
            (v6 / "figure2_final_external_rescue.png", 6.3),
            (v6 / "figure3_final_subject_confusion.png", 5.6),
        ],
        5: [(v6 / "figure5_final_error_profiles.png", 6.3)],
        6: [
            (algo / "figure_algorithm_ablation.png", 6.0),
            (algo / "figure_calibration_reliability.png", 5.8),
            (algo / "figure_risk_constraint_curve.png", 5.8),
        ],
        7: [
            (brain / "figure_B5_papez_network.png", 6.2),
            (brain / "figure_B1_anatomy_underlay_montage.png", 6.3),
        ],
    }


def appendix_figures(root: Path) -> list[tuple[str, Path, float]]:
    brain = root / "reports/brain_figures_fcstyle/figures"
    return [
        ("Supplementary Figure S1. Anatomy-underlaid method montage (supporting spatial maps).", brain / "figure_B1_anatomy_underlay_montage.png", 6.3),
        ("Supplementary Figure S2. ROI × method/stage heatmaps.", brain / "figure_B2_roi_method_heatmap_panels.png", 6.3),
        ("Supplementary Figure S3. Top-20 ROI multi-method overlay.", brain / "figure_B3_top_roi_dots.png", 6.3),
        ("Supplementary Figure S4. DKT cortical surface atrophy (AD−CN).", brain / "figure_B4b_dkt_surface.png", 6.0),
        ("Supplementary Figure S5. Glass-brain summary maps.", brain / "figure_B4_glassbrain_summary.png", 6.0),
        ("Supplementary Figure S6. Subject-level MRF-q case montages.", brain / "figure_B6_case_montages.png", 6.2),
        ("Supplementary Figure S7. OASIS stress-test spatial cases (boundary/limitation).", brain / "figure_B7_oasis_stress.png", 5.5),
    ]


def parse_captions(lines: list[str]) -> dict[int, str]:
    caps: dict[int, str] = {}
    for line in lines:
        m = re.match(r"^\*\*Figure\s+(\d+)\.\s*(.*?)\*\*\s*(.*)$", line.strip())
        if not m:
            continue
        n = int(m.group(1))
        title = m.group(2).strip()
        rest = m.group(3).strip()
        caps[n] = f"Figure {n}. {title} {rest}".strip()
    return caps


def latex_to_display(expr: str) -> str:
    """Normalize display LaTeX for matplotlib mathtext."""
    text = " ".join(expr.strip().split())
    text = text.replace("\\qquad", r"\quad")
    text = text.replace("\\mathrm{CN}", r"\mathrm{CN}")
    text = text.replace("\\mathrm{MCI}", r"\mathrm{MCI}")
    text = text.replace("\\mathrm{AD}", r"\mathrm{AD}")
    # mathtext does not like \max as operatorname sometimes; keep \max
    return text


def render_equation_png(latex: str, cache_dir: Path, eq_no: int) -> Path:
    cache_dir.mkdir(parents=True, exist_ok=True)
    key = hashlib.md5(f"{eq_no}|{latex}".encode()).hexdigest()[:12]
    out = cache_dir / f"eq_{eq_no}_{key}.png"
    if out.exists():
        return out

    display = latex_to_display(latex)
    fig = plt.figure(figsize=(6.8, 0.95), dpi=220)
    fig.patch.set_facecolor("white")
    ax = fig.add_axes([0.02, 0.05, 0.86, 0.9])
    ax.axis("off")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    try:
        ax.text(0.5, 0.5, f"${display}$", ha="center", va="center", fontsize=13, color="#1a1a1a")
    except Exception:
        # Fallback: plain text
        ax.text(0.5, 0.5, display, ha="center", va="center", fontsize=11, color="#1a1a1a")
    ax_num = fig.add_axes([0.88, 0.05, 0.10, 0.9])
    ax_num.axis("off")
    ax_num.text(0.5, 0.5, f"({eq_no})", ha="center", va="center", fontsize=12, color="#333333")
    fig.savefig(out, bbox_inches="tight", pad_inches=0.12, facecolor="white")
    plt.close(fig)
    return out


def unicode_inline_math(expr: str) -> str:
    text = expr.strip()
    repl = [
        (r"\tilde{p}", "p̃"),
        (r"\bar{p}", "p̄"),
        (r"\hat{y}", "ŷ"),
        (r"\mathrm{CN}", "CN"),
        (r"\mathrm{MCI}", "MCI"),
        (r"\mathrm{AD}", "AD"),
        (r"\epsilon", "ε"),
        (r"\ge", "≥"),
        (r"\le", "≤"),
        (r"\in", "∈"),
        (r"\sum", "Σ"),
        (r"\{", "{"),
        (r"\}", "}"),
    ]
    for old, new in repl:
        text = text.replace(old, new)
    text = re.sub(r"_\{([^}]+)\}", r"_\1", text)
    text = re.sub(r"\^\{([^}]+)\}", r"^\1", text)
    text = text.replace("\\", "")
    return text


def prepare_body_text(text: str) -> str:
    """Convert inline $...$ / \\(...\\) math to readable Unicode-ish text."""
    text = re.sub(r"\\\((.+?)\\\)", lambda m: unicode_inline_math(m.group(1)), text)
    text = re.sub(r"(?<!\$)\$(?!\$)(.+?)(?<!\$)\$(?!\$)", lambda m: unicode_inline_math(m.group(1)), text)
    return text


def add_caption(doc: Document, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=2, after=12, line=1.10)
    # Bold "Figure N. title" prefix if present
    m = re.match(r"^(Figure\s+\d+\.\s+[^.]+?\.)\s*(.*)$", caption)
    if m:
        run = p.add_run(m.group(1) + " ")
        set_run_font(run, size=9.5, bold=True, italic=True)
        if m.group(2):
            run2 = p.add_run(m.group(2))
            set_run_font(run2, size=9.5, italic=True)
    else:
        run = p.add_run(caption)
        set_run_font(run, size=9.5, italic=True)


def add_image(doc: Document, path: Path, width: float) -> None:
    if not path.exists():
        p = add_paragraph(doc, f"[Missing figure: {path.name}]", before=6, after=4, base_size=10)
        for run in p.runs:
            run.font.italic = True
            run.font.color.rgb = GRAY
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, before=8, after=2, line=1.05)
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))


def insert_figure_block(
    doc: Document,
    fig_n: int,
    assets: dict[int, list[tuple[Path, float]]],
    captions: dict[int, str],
) -> None:
    paths = assets.get(fig_n, [])
    for path, width in paths:
        add_image(doc, path, width)
    caption = captions.get(fig_n, f"Figure {fig_n}.")
    add_caption(doc, caption)


def insert_equation(doc: Document, latex: str, cache_dir: Path, eq_no: int) -> None:
    png = render_equation_png(latex, cache_dir, eq_no)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, before=6, after=8, line=1.05)
    run = paragraph.add_run()
    run.add_picture(str(png), width=Inches(6.2))


def build_docx(source: Path, output: Path, root: Path) -> None:
    doc = Document()
    setup_styles(doc)
    lines = source.read_text(encoding="utf-8").splitlines()
    captions = parse_captions(lines)
    assets = figure_assets(root)
    eq_cache = root / "reports/v6_final_model/equation_cache"
    eq_no = 0
    placed_figs: set[int] = set()

    i = 0
    in_math = False
    math_buffer: list[str] = []
    in_caption_section = False
    skip_caption_bodies = False

    while i < len(lines):
        raw = lines[i]
        text = raw.strip()
        if not text:
            i += 1
            continue

        # Inline figure marker
        m_fig = re.fullmatch(r"\{\{FIG:(\d+)\}\}", text)
        if m_fig:
            fig_n = int(m_fig.group(1))
            insert_figure_block(doc, fig_n, assets, captions)
            placed_figs.add(fig_n)
            i += 1
            continue

        if text == "\\[":
            in_math = True
            math_buffer = []
            i += 1
            continue
        if in_math:
            if text == "\\]":
                eq_no += 1
                insert_equation(doc, " ".join(math_buffer), eq_cache, eq_no)
                in_math = False
            else:
                math_buffer.append(text)
            i += 1
            continue

        if text.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        if text.startswith("## Figure captions"):
            in_caption_section = True
            # Keep a compact caption list without re-embedding images
            doc.add_page_break()
            add_paragraph(doc, "Figure captions (reference list)", style="Heading 1")
            add_paragraph(
                doc,
                "Figures 1–7 are embedded inline in Methods/Results in evidence-chain order. "
                "This section repeats captions for journal-style reference.",
                after=8,
                base_size=10.5,
            )
            i += 1
            continue

        if in_caption_section and text.startswith("**Figure"):
            # already shown inline; keep text-only reference
            add_paragraph(doc, strip_inline_md(text), before=4, after=8, base_size=9.5)
            if doc.paragraphs:
                for run in doc.paragraphs[-1].runs:
                    run.font.italic = True
            i += 1
            continue

        if text.startswith("# "):
            title = strip_inline_md(text[2:].strip())
            paragraph = doc.add_paragraph()
            set_paragraph_spacing(paragraph, before=0, after=3, line=1.05)
            run = paragraph.add_run(title)
            set_run_font(run, size=20, bold=True, color=DARK_BLUE)
            subtitle = doc.add_paragraph()
            set_paragraph_spacing(subtitle, before=0, after=12, line=1.10)
            run = subtitle.add_run(
                "V6 manuscript with inline figures and rendered equations "
                "(research prototype, not for clinical use)"
            )
            set_run_font(run, size=10.5, color=GRAY)
        elif text.startswith("## "):
            # Leaving caption section for later headings
            if in_caption_section and not text.startswith("## Figure"):
                in_caption_section = False
            add_paragraph(doc, strip_inline_md(text[3:].strip()), style="Heading 1")
        elif text.startswith("### "):
            add_paragraph(doc, strip_inline_md(text[4:].strip()), style="Heading 2")
        elif re.match(r"^\d+\.\s+", text):
            item = prepare_body_text(re.sub(r"^\d+\.\s+", "", text))
            paragraph = add_paragraph(doc, item, after=4, line=1.10)
            paragraph.style = "List Number"
        elif text.startswith("- "):
            paragraph = add_paragraph(doc, prepare_body_text(text[2:]), after=4, line=1.10)
            paragraph.style = "List Bullet"
        else:
            add_paragraph(doc, prepare_body_text(text), after=6, line=1.10)
        i += 1

    # Any missing inline figures (safety)
    missing = [n for n in range(1, 8) if n not in placed_figs]
    if missing:
        doc.add_page_break()
        add_paragraph(doc, "Additional figures (not placed inline)", style="Heading 1")
        for n in missing:
            insert_figure_block(doc, n, assets, captions)

    # Supplementary neuroimaging
    doc.add_page_break()
    add_paragraph(doc, "Supplementary neuroimaging figures", style="Heading 1")
    add_paragraph(
        doc,
        "Supporting spatial panels for the anatomical-consistency branch. "
        "These are atlas-filled or DKT-proxy maps, not voxel-wise FDR discoveries.",
        after=10,
        base_size=10.5,
    )
    for caption, path, width in appendix_figures(root):
        add_image(doc, path, width)
        add_caption(doc, caption)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    print(f"[saved] {output}")
    print(f"[inline figures] {sorted(placed_figs)}")
    print(f"[equations] {eq_no}")
    print(f"[appendix] {len(appendix_figures(root))}")


def main() -> None:
    root = resolve_repo_root()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source",
        type=Path,
        default=root / "reports/v6_final_model/manuscript_v6_full_draft.md",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=root / "reports/v6_final_model/ARA-Net_V6_manuscript_with_figures.docx",
    )
    parser.add_argument("--repo-root", type=Path, default=root)
    args = parser.parse_args()
    build_docx(args.source, args.output, args.repo_root.resolve())


if __name__ == "__main__":
    main()
