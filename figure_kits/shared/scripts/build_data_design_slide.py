#!/usr/bin/env python3
"""Recreate the user's 'Data Organization & Experimental Design / Chapter-wise
Tasks and Evaluation' figure as a single fully-editable widescreen PPT slide
(native shapes + text, serif font to match the source). Text and structure
mirror the source figure; fitted to 16:9.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "Data-Organization-Experimental-Design-EN.pptx"

FONT = "Times New Roman"

NAVY = RGBColor(0x1B, 0x2A, 0x44)
DARK = RGBColor(0x20, 0x24, 0x2C)
GRAY = RGBColor(0x55, 0x5B, 0x66)
BORDER = RGBColor(0x8A, 0x8F, 0x99)
PANEL_BORDER = RGBColor(0xB9, 0xBE, 0xC8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HDR_BLUE = RGBColor(0xDC, 0xE6, 0xF2)
TAG_BLUE = RGBColor(0xC7, 0xD6, 0xEC)
HDR_GREEN = RGBColor(0xDD, 0xEA, 0xD6)
YELLOW = RGBColor(0xFB, 0xF2, 0xD3)
ARROW = RGBColor(0x3A, 0x4C, 0x6B)

SLIDE_W, SLIDE_H = 13.333, 7.5


def set_dash(shape, val="dash"):
    ln = shape.line._get_or_add_ln()
    for e in ln.findall(qn("a:prstDash")):
        ln.remove(e)
    d = ln.makeelement(qn("a:prstDash"), {"val": val})
    ln.append(d)


def rect(slide, x, y, w, h, fill, line=None, line_w=1.0, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    return sp


def rich_text(slide, x, y, w, h, paras, *, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE, pad=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", align)
        para.space_after = Pt(p.get("sa", 1))
        para.space_before = Pt(p.get("sb", 0))
        para.line_spacing = p.get("ls", 1.0)
        run = para.add_run()
        run.text = p["t"]
        f = run.font
        f.name = p.get("font", FONT)
        f.size = Pt(p["s"])
        f.bold = p.get("b", False)
        f.italic = p.get("i", False)
        f.color.rgb = p.get("c", DARK)
    return box


def t(text, *, b=False, c=DARK, s=8.0, i=False, align=None, sa=1, ls=1.0):
    d = {"t": text, "b": b, "c": c, "s": s, "i": i, "sa": sa, "ls": ls}
    if align is not None:
        d["align"] = align
    return d


def header_bar(slide, x, y, w, h, text, fill, *, size=10.5):
    rect(slide, x, y, w, h, fill, line=BORDER, line_w=0.75)
    rich_text(slide, x, y, w, h, [t(text, b=True, c=NAVY, s=size)])


def box_text(slide, x, y, w, h, paras, *, fill=WHITE, border=BORDER, dash=False,
             line_w=1.0, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER):
    sp = rect(slide, x, y, w, h, fill, line=border, line_w=line_w)
    if dash:
        set_dash(sp)
    rich_text(slide, x, y, w, h, paras, anchor=anchor, align=align)


def arrow(slide, x, y, w, h, *, shape=MSO_SHAPE.RIGHT_ARROW, color=ARROW):
    return rect(slide, x, y, w, h, color, shape=shape)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # =============================== LEFT PANEL =============================== #
    LX, LY, LW, LH = 0.18, 0.16, 5.55, 7.20
    rect(slide, LX, LY, LW, LH, None, line=PANEL_BORDER, line_w=1.25)
    rich_text(slide, LX, LY + 0.06, LW, 0.36,
              [t("DATA ORGANIZATION AND EXPERIMENTAL DESIGN", b=True, c=NAVY, s=12)])

    # two columns: A (data sources) | B (diagnostic labels)
    AX, AW = 0.34, 2.62
    BX, BW = 3.06, 2.55

    header_bar(slide, AX, 0.60, AW, 0.30, "A. Data Sources and Roles", HDR_BLUE)
    header_bar(slide, BX, 0.60, BW, 0.30, "B. Diagnostic Labels", HDR_GREEN)

    # --- A. four dataset rows ---
    def datarow(x, y, w, h, name, role, desc):
        rect(slide, x, y, w, h, WHITE, line=BORDER, line_w=0.9)
        nw = 0.82
        rich_text(slide, x + 0.02, y, nw, h, [t(name, b=True, c=NAVY, s=12.5)])
        line = rect(slide, x + nw, y + 0.08, 0.008, h - 0.16, BORDER)
        rich_text(slide, x + nw + 0.06, y, w - nw - 0.1, h,
                  [t(role, b=True, c=NAVY, s=8.4, sa=1),
                   t(desc, c=GRAY, s=7.0, ls=1.0)],
                  align=PP_ALIGN.CENTER)

    ay, ah, ag = 0.98, 0.92, 0.06
    datarow(AX, ay, AW, ah, "ADNI", "Development & Internal Testing Set",
            "Model training, hyperparameter tuning, and internal validation")
    datarow(AX, ay + (ah + ag), AW, ah, "AIBL", "Independent External Validation Set",
            "Locked after model selection; one-time held-out test to evaluate "
            "cross-site/protocol generalization")
    datarow(AX, ay + 2 * (ah + ag), AW, ah, "IXI", "Healthy Control Set",
            "Includes only cognitively normal (CN) subjects; used to assess "
            "specificity (healthy subjects not misdiagnosed)")
    datarow(AX, ay + 3 * (ah + ag), AW, ah, "OASIS", "Stress Test Set",
            "Zero-shot evaluation (no fine-tuning) to expose the limits of "
            "domain transfer")

    # --- B. diagnostic-label boxes ---
    box_text(slide, BX, 0.98, BW, 1.02,
             [t("Primary Three-Class Groups", b=True, c=NAVY, s=8.6),
              t("CN (Cognitively Normal) /", c=DARK, s=7.6),
              t("MCI (Mild Cognitive Impairment) /", c=DARK, s=7.6),
              t("AD (Alzheimer's Disease)", c=DARK, s=7.6)])
    box_text(slide, BX, 2.12, BW, 1.30,
             [t("MCI Subtyping (Longitudinal Outcome)", b=True, c=NAVY, s=8.4),
              t("pMCI (Progressive MCI)", c=DARK, s=7.6),
              t("\u2192 converts to AD", c=GRAY, s=7.4),
              t("sMCI (Stable MCI)", c=DARK, s=7.6),
              t("\u2192 remains stable", c=GRAY, s=7.4)],
             dash=True)
    box_text(slide, BX, 3.54, BW, 1.00,
             [t("Chapter 4 Binary Setting", b=True, c=NAVY, s=8.6),
              t("AD vs NC (Normal Control)", c=DARK, s=7.8),
              t("(MCI subjects excluded", c=GRAY, s=7.2),
              t("in the main experiments)", c=GRAY, s=7.2)],
             fill=YELLOW)

    # --- C. data handling principles (full width) ---
    header_bar(slide, AX, 5.00, BX + BW - AX, 0.30, "C. Data Handling Principles", HDR_BLUE)
    box_text(slide, AX, 5.38, BX + BW - AX, 0.86,
             [t("All splits are subject-level (by participant, not by scan)", b=True, c=NAVY, s=8.8),
              t("to prevent information leakage from multiple longitudinal scans of the "
                "same participant across training/test sets.", c=GRAY, s=7.6)])
    box_text(slide, AX, 6.32, BX + BW - AX, 0.80,
             [t("Imaging Modalities", b=True, c=NAVY, s=8.8),
              t("T1-weighted MRI, FLAIR MRI (used individually or in combination "
                "depending on chapter-specific settings)", c=GRAY, s=7.6)])

    # A -> B arrow
    arrow(slide, AX + AW + 0.02, 1.30, BX - (AX + AW) - 0.04, 0.20)

    # =============================== RIGHT PANEL ============================== #
    RX, RY, RW, RH = 5.92, 0.16, 7.26, 7.20
    rect(slide, RX, RY, RW, RH, None, line=PANEL_BORDER, line_w=1.25)
    rich_text(slide, RX, RY + 0.06, RW, 0.36,
              [t("CHAPTER-WISE TASKS AND EVALUATION", b=True, c=NAVY, s=12)])

    inner_x = RX + 0.16
    inner_w = RW - 0.32
    tag_w, task_w, cgap = 1.50, 2.30, 0.10
    data_w = inner_w - tag_w - task_w - 2 * cgap
    task_x = inner_x + tag_w + cgap
    data_x = task_x + task_w + cgap

    chapters = [
        ("1", "ARA-Net",
         ["Multi-class classification:", "CN / MCI / AD"],
         ["AIBL subject-level", "external test set"]),
        ("2", "PathwayPro",
         ["Multi-class classification:", "CN / MCI / AD"],
         ["ADNI paired T1+FLAIR subset for model training and internal "
          "validation; OASIS for zero-shot (no fine-tuning) test"]),
        ("3", "A2C-NODE",
         ["Multi-class classification: CN / MCI / AD", "MCI subtyping: pMCI vs sMCI"],
         ["ADNI longitudinal internal test set for development; AIBL external "
          "test set for generalization; Analysis on pMCI / sMCI subgroups"]),
        ("4", "CUED-AD",
         ["Binary classification:", "AD vs NC", "(MCI excluded)"],
         ["AIBL external conflict pool including natural conflicts and "
          "non-conflict samples"]),
    ]

    ry0, rh, rgap = 0.64, 1.22, 0.10
    for k, (num, name, task_lines, data_lines) in enumerate(chapters):
        oy = ry0 + k * (rh + rgap)
        outer = rect(slide, inner_x, oy, inner_w, rh, None, line=BORDER, line_w=1.0,
                     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        set_dash(outer)
        pad = 0.10
        cy, ch = oy + pad, rh - 2 * pad
        box_text(slide, inner_x + pad, cy, tag_w, ch,
                 [t(f"Chapter {num}", b=True, c=NAVY, s=10),
                  t(name, b=True, c=NAVY, s=11)],
                 fill=TAG_BLUE)
        box_text(slide, task_x, cy, task_w, ch,
                 [t("Task", b=True, c=NAVY, s=8.6)] +
                 [t(l, c=DARK, s=7.6) for l in task_lines])
        box_text(slide, data_x, cy, data_w, ch,
                 [t("Data & Evaluation", b=True, c=NAVY, s=8.6)] +
                 [t(l, c=DARK, s=7.4) for l in data_lines])
        # arrow from left panel into the chapter tag
        arrow(slide, RX - 0.30, oy + rh / 2 - 0.10, 0.36, 0.20)

    # Outputs box (spanning) + arrow from left
    oy = ry0 + len(chapters) * (rh + rgap)
    box_text(slide, inner_x, oy + 0.02, inner_w, RY + RH - (oy + 0.02) - 0.12,
             [t("Outputs", b=True, c=NAVY, s=10.5),
              t("Performance metrics (e.g., Accuracy, AUC, Sensitivity, Specificity), "
                "cross-dataset generalization analysis and transfer boundary assessment",
                c=DARK, s=8.0)],
             fill=HDR_BLUE)
    arrow(slide, RX - 0.30, oy + 0.20, 0.36, 0.20)

    prs.save(str(OUTPUT))
    print(f"Saved: {OUTPUT}")
    print(f"Slides: {len(prs.slides)} | shapes: {len(slide.shapes)}")
    return OUTPUT


if __name__ == "__main__":
    build()
