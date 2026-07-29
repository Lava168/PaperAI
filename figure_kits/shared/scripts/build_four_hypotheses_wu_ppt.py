#!/usr/bin/env python3
"""Build 4 English hypothesis slides using the Wu group-meeting deck background style.

Each slide pairs an open-access / project literature figure with the corresponding
chapter architecture figure, plus compact Problem / Hypothesis / Methods / Evidence cards.

Template reference: 20241107 吴奕佳 大组会汇报.pptx
"""

from __future__ import annotations

import shutil
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = Path(__file__).resolve().parent
WU_TEMPLATE = ROOT / "20241107 吴奕佳 大组会汇报.pptx"
OUTPUT = ROOT / "Four-Core-Problems-Scientific-Hypotheses-EN-WuStyle.pptx"

ASSETS = ROOT / "ppt_assets_extracted" / "ppt_images-2"
LIT = SCRIPTS / "_lit_figures"

# Wu deck palette
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BADGE = RGBColor(0x44, 0x72, 0xC4)
BAR = RGBColor(0x00, 0x70, 0xC0)
DARK = RGBColor(0x20, 0x20, 0x20)
GRAY = RGBColor(0x66, 0x66, 0x66)
BLUE_TXT = RGBColor(0x00, 0x56, 0xA0)
GREEN = RGBColor(0x00, 0x70, 0x50)
ACCENT = RGBColor(0xC0, 0x55, 0x10)
LIGHT_BG = RGBColor(0xF2, 0xF6, 0xFB)
BORDER = RGBColor(0xC5, 0xD5, 0xE8)

SLIDE_W = 13.333
SLIDE_H = 7.5
FONT = "Arial"


def _no_line(shape):
    shape.line.fill.background()


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    _no_line(shape)


def delete_all_slides(prs: Presentation) -> None:
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]


def first_existing(*paths: Path) -> Path | None:
    for p in paths:
        if p and p.exists():
            return p
    return None


def fit_image(path: Path, max_w: float, max_h: float) -> tuple[float, float]:
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih
    w, h = max_w, max_w / ar
    if h > max_h:
        h = max_h
        w = h * ar
    return w, h


def path_for_pptx_embed(img_path: Path) -> Path:
    """Return a path python-pptx can embed (convert WEBP etc. to PNG)."""
    if img_path.suffix.lower() in {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff"}:
        return img_path
    cache_dir = SCRIPTS / "_pptx_embed_cache"
    cache_dir.mkdir(exist_ok=True)
    out = cache_dir / f"{img_path.name}.png"
    if not out.exists() or out.stat().st_mtime < img_path.stat().st_mtime:
        with Image.open(img_path) as im:
            im.convert("RGB").save(out, "PNG")
    return out


def textbox(slide, left, top, width, height, text, *, size=11, color=DARK,
            bold=False, italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return box


def rich_block(slide, left, top, width, height, lines, *, size=9.5, line_spacing=1.06):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    first = True
    for txt, bold, color in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(3)
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = txt
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def wu_header(slide, number: str, section: str, title: str):
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Inches(-0.006), Inches(0.273), Inches(0.53), Inches(0.54))
    badge.adjustments[0] = 0.18
    _fill(badge, BADGE)
    textbox(slide, -0.006, 0.273, 0.53, 0.54, number, size=20, color=WHITE, bold=True,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(slide, 0.71, 0.29, 2.2, 0.50, section, size=18, color=DARK, bold=True,
            anchor=MSO_ANCHOR.MIDDLE)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                 Inches(2.70), Inches(0.35), Inches(10.40), Inches(0.42))
    _fill(bar, BAR)
    textbox(slide, 2.72, 0.30, 10.0, 0.50, title, size=16.5, color=WHITE, bold=True,
            anchor=MSO_ANCHOR.MIDDLE)


def footer_strip(slide, page_num: int):
    sq, gap, x0, y0 = 0.18, 0.04, 0.13, 7.28
    for i in range(5):
        s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Inches(x0 + i * (sq + gap)), Inches(y0), Inches(sq), Inches(sq))
        s.adjustments[0] = 0.12
        _fill(s, BADGE if i < 4 else RGBColor(0x5B, 0x9B, 0xD5))
    textbox(slide, 6.15, 7.12, 1.0, 0.35, str(page_num), size=11, color=GRAY,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def card(slide, x, y, w, h, title, lines, *, title_color=BLUE_TXT, size=9.3):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    panel.adjustments[0] = 0.04
    panel.fill.solid()
    panel.fill.fore_color.rgb = LIGHT_BG
    panel.line.color.rgb = BORDER
    panel.line.width = Pt(0.75)
    textbox(slide, x + 0.10, y + 0.06, w - 0.20, 0.24, title, size=10.5, color=title_color, bold=True)
    rich_block(slide, x + 0.10, y + 0.30, w - 0.20, h - 0.36, lines, size=size)


def image_panel(slide, x, y, w, h, img_path: Path, label: str, citation: str):
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    frame.adjustments[0] = 0.03
    frame.fill.solid()
    frame.fill.fore_color.rgb = WHITE
    frame.line.color.rgb = BORDER
    frame.line.width = Pt(0.75)

    cap_h, cite_h = 0.22, 0.34
    pad = 0.08
    img_box_h = h - cap_h - cite_h - pad * 2
    iw, ih = fit_image(img_path, w - pad * 2, img_box_h)
    ix = x + (w - iw) / 2
    iy = y + pad + (img_box_h - ih) / 2
    slide.shapes.add_picture(str(path_for_pptx_embed(img_path)), Inches(ix), Inches(iy), Inches(iw), Inches(ih))

    textbox(slide, x + 0.08, y + h - cap_h - cite_h - 0.02, w - 0.16, cap_h,
            label, size=9.5, color=BLUE_TXT, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, x + 0.08, y + h - cite_h - 0.02, w - 0.16, cite_h,
            citation, size=7.5, color=GRAY, italic=True, align=PP_ALIGN.CENTER)


SLIDES = [
    dict(
        num="1",
        section="Core Problem I",
        title="Interpretability Gap  →  ARA-Net (Atlas-Guided Multimodal Staging)",
        lit_img=first_existing(
            LIT / "npj_s41746-024-01123-7_Fig1.png",
            ASSETS / "bg_explainable_AI_AD_SECNN.webp",
        ),
        lit_label="Literature: ante-hoc XAI pipeline for dementia MRI",
        lit_cite="Leonardsen EH et al. npj Digit Med. 2024;7:123 (CC BY 4.0)",
        ours_img=first_existing(
            ASSETS / "slide9_ARANet_architecture.png",
            ROOT / "chapter1_foundation/ARA-Net/reports/v6_final_model/figures/figure3_rcspe_architecture_nbe_style.png",
        ),
        ours_label="Our method: ARA-Net atlas-guided RC-SPE staging",
        ours_cite="Zhao Y et al. ARA-Net v6 manuscript. Tsinghua SIGS; 2026",
        problem=[
            ("Core problem:", True, ACCENT),
            ("Many AD models output labels only; post-hoc saliency ≠ valid biological explanation.", False, DARK),
        ],
        hypothesis=[
            ("Scientific hypothesis (H1):", True, BLUE_TXT),
            ("21-region atlas staging + RC-SPE yields external validation AND auditable "
             "neurodegeneration consistency (not attention = biomarker).", False, DARK),
        ],
        methods=[
            ("Methods:", True, DARK),
            ("• 21-region atlas + clinical vars → 6-stream RC-SPE", False, DARK),
            ("• AD-key consistency + permutation test", False, DARK),
        ],
        evidence=[
            ("Evidence:", True, GREEN),
            ("• AIBL BAcc 0.833 | Macro-AUC 0.937", False, DARK),
            ("• AD-key 0.510 vs null 0.286 (p=0.026)", False, DARK),
            ("Limitation: MCI recall 0.686.", True, ACCENT),
        ],
    ),
    dict(
        num="2",
        section="Core Problem II",
        title="Multi-Mechanism Gap  →  PathwayPro (Atrophy vs Vascular Disentanglement)",
        lit_img=first_existing(
            LIT / "PMC9924910_Figure1_from_pdfimg_p3_1.png",
            LIT / "PMC9924910_Figure1_fixel_illustration_crop.png",
            LIT / "PMC9924910_Figure1_from_pdfimg_p3_0.png",
            ASSETS / "slide10_disentangled_imaging.jpg",
        ),
        lit_label="Literature: disentangling AD vs small-vessel disease on WM",
        lit_cite="Franzmeier N et al. Brain. 2022;145:3571 (PMC9924910, CC BY-NC)",
        ours_img=first_existing(
            ASSETS / "slide10_PathwayPro_architecture.png",
            ROOT / "chapter2_disentangle/paper/figures/fig1_pipeline.png",
        ),
        ours_label="Our method: PathwayPro T1/FLAIR dual-pathway disentanglement",
        ours_cite="Zhao Y et al. PathwayPro manuscript. Tsinghua SIGS; 2026",
        problem=[
            ("Core problem:", True, ACCENT),
            (">60% AD has vascular comorbidity, yet models collapse atrophy & vascular signals.", False, DARK),
        ],
        hypothesis=[
            ("Scientific hypothesis (H2):", True, BLUE_TXT),
            ("Dual encoders + HSIC/TC/CLUB/iVAE separate z_a (atrophy) from z_v (vascular/WMH).", False, DARK),
        ],
        methods=[
            ("Methods:", True, DARK),
            ("• Dual 3D ResNet + modality routing + cross-attention", False, DARK),
            ("• Zero-mask T1+FLAIR / T1-only deployment", False, DARK),
        ],
        evidence=[
            ("Evidence:", True, GREEN),
            ("• z_v→WMH AUC 0.923 | OASIS zero-shot AUC 0.774", False, DARK),
            ("Limitation: disent score −0.014; AIBL zero-shot fails.", True, ACCENT),
        ],
    ),
    dict(
        num="3",
        section="Core Problem III",
        title="Dynamics Gap  →  A2C-NODE (Continuous-Time Causal Dynamics)",
        lit_img=first_existing(
            LIT / "PMC12407814_Figure1.png",
            LIT / "gnova_page-02.png",
            ASSETS / "slide11_latent_ODE.png",
        ),
        lit_label="Literature: NeuralODE DPM for sparse multimodal AD progression",
        lit_cite="Zanin A et al. bioRxiv 2025 (PMC12407814; CC BY 4.0)",
        ours_img=first_existing(
            ASSETS / "slide11_A2C_NODE_architecture.png",
            ROOT / "chapter3_neural_ode/A2C-NODE/overleaf_a2c_media/figures/Figure1_A2C_NODE_framework.png",
        ),
        ours_label="Our method: 21-region graph neural ODE + ATE report",
        ours_cite="Zhao Y et al. A2C-NODE manuscript. Tsinghua SIGS; 2026",
        problem=[
            ("Core problem:", True, ACCENT),
            ("Static snapshots; LSTM needs fixed grids and struggles with irregular follow-up.", False, DARK),
        ],
        hypothesis=[
            ("Scientific hypothesis (H3):", True, BLUE_TXT),
            ("Graph neural ODE captures irregular trajectories better than LSTM; "
             "region-clamp ATE hints at causal drivers.", False, DARK),
        ],
        methods=[
            ("Methods:", True, DARK),
            ("• Hybrid graph A = α·A_prior + (1−α)·attention", False, DARK),
            ("• GCN Neural ODE + multi-task readout + ATE", False, DARK),
        ],
        evidence=[
            ("Evidence:", True, GREEN),
            ("• BAcc 0.854 vs LSTM 0.678 (+17.6 pp)", False, DARK),
            ("Limitation: AIBL raw Acc 0.297 (cohort shift).", True, ACCENT),
        ],
    ),
    dict(
        num="4",
        section="Core Problem IV",
        title="Trust Gap  →  CUED-AD (Conflict- & Uncertainty-Aware Decision Support)",
        lit_img=first_existing(
            LIT / "jin2024_npj_hidden_flaws_fig1.png",
            ASSETS / "slide12_UQ_review.jpg",
        ),
        lit_label="Literature: high accuracy can hide flawed GPT-4V rationales",
        lit_cite="Jin Q et al. npj Digit Med. 2024;7:190 (CC BY 4.0)",
        ours_img=first_existing(
            ASSETS / "slide12_CUED_AD_architecture.png",
            ROOT / "chapter4_generalization/demo_assets/images/cued_ad_system_overview.png",
        ),
        ours_label="Our method: CUED-AD conflict + uncertainty decision support",
        ours_cite="Zhao Y et al. CUED-AD manuscript. Tsinghua SIGS; 2026",
        problem=[
            ("Core problem:", True, ACCENT),
            ("High lab AUC ≠ clinical trust; missing uncertainty, conflict, and review tiers.", False, DARK),
        ],
        hypothesis=[
            ("Scientific hypothesis (H4):", True, BLUE_TXT),
            ("D_JS conflict + U_total cluster in natural-conflict cases → escalate to human review.", False, DARK),
        ],
        methods=[
            ("Methods:", True, DARK),
            ("• Dual MC-Dropout branches + SCM residuals", False, DARK),
            ("• D_JS + U_total tiered review prototype", False, DARK),
        ],
        evidence=[
            ("Evidence:", True, GREEN),
            ("• AIBL Acc 0.984 | conflict AUC ~0.75", False, DARK),
            ("Limitation: binary task; research prototype only.", True, ACCENT),
        ],
    ),
]


def build_slide(prs, data: dict, page: int):
    if not data.get("lit_img") or not data.get("ours_img"):
        missing = []
        if not data.get("lit_img"):
            missing.append("literature image")
        if not data.get("ours_img"):
            missing.append("architecture image")
        raise SystemExit(f"Slide {page}: missing {', '.join(missing)}")

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    wu_header(slide, data["num"], data["section"], data["title"])

    margin_l, margin_r = 0.55, 0.55
    gap = 0.18
    total_w = SLIDE_W - margin_l - margin_r
    col_w = (total_w - gap) / 2

    # Row 1 — compact text
    top1 = 1.02
    row1_h = 1.18
    card(slide, margin_l, top1, col_w, row1_h, "Problem", data["problem"], title_color=ACCENT, size=9.0)
    card(slide, margin_l + col_w + gap, top1, col_w, row1_h, "Hypothesis", data["hypothesis"], size=9.0)

    # Row 2 — literature + our architecture figures
    top2 = top1 + row1_h + 0.14
    row2_h = 2.72
    image_panel(slide, margin_l, top2, col_w, row2_h,
                data["lit_img"], data["lit_label"], data["lit_cite"])
    image_panel(slide, margin_l + col_w + gap, top2, col_w, row2_h,
                data["ours_img"], data["ours_label"], data["ours_cite"])

    # Row 3 — methods + evidence
    top3 = top2 + row2_h + 0.14
    row3_h = 1.42
    card(slide, margin_l, top3, col_w, row3_h, "Methods", data["methods"], title_color=DARK, size=9.0)
    card(slide, margin_l + col_w + gap, top3, col_w, row3_h,
         "Validation & Evidence", data["evidence"], title_color=GREEN, size=9.0)

    footer_strip(slide, page)
    return slide


def main():
    if not WU_TEMPLATE.exists():
        raise SystemExit(f"Wu template not found: {WU_TEMPLATE}")

    tmp = OUTPUT.with_suffix(".tmp.pptx")
    shutil.copy2(WU_TEMPLATE, tmp)
    prs = Presentation(str(tmp))
    delete_all_slides(prs)

    for i, data in enumerate(SLIDES, start=1):
        build_slide(prs, data, i)

    prs.save(str(tmp))
    tmp.replace(OUTPUT)
    print(f"Saved: {OUTPUT}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
