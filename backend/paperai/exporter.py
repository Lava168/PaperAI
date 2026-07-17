from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


def combined_markdown(project: dict, artifacts: list[dict]) -> str:
    final = next((item for item in artifacts if item["step"] == "final_manuscript"), None)
    if final:
        return final["content"]
    blocks = [f"# {project['title']}", ""]
    for artifact in artifacts:
        if artifact["step"].startswith("feedback_"):
            continue
        blocks.extend([f"## {artifact['step'].replace('_', ' ').title()}", "", artifact["content"], ""])
    return "\n".join(blocks)


def markdown_to_docx(markdown: str, output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped:
            document.add_paragraph()
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            document.add_heading(stripped[2:], level=1)
        elif re.match(r"^[-*] ", stripped):
            document.add_paragraph(stripped[2:], style="List Bullet")
        elif re.match(r"^\d+\. ", stripped):
            document.add_paragraph(re.sub(r"^\d+\. ", "", stripped), style="List Number")
        else:
            paragraph = document.add_paragraph()
            _add_inline(paragraph, stripped)
    document.save(output)


def _add_inline(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
        else:
            paragraph.add_run(part)


def markdown_to_latex(markdown: str) -> str:
    escaped = markdown.replace("\\", r"\textbackslash{}")
    for source, target in [("&", r"\&"), ("%", r"\%"), ("$", r"\$"), ("#", r"\#"), ("_", r"\_")]:
        escaped = escaped.replace(source, target)
    lines = []
    for line in escaped.splitlines():
        if line.startswith(r"\#\#\# "):
            lines.append(r"\subsubsection{" + line[7:] + "}")
        elif line.startswith(r"\#\# "):
            lines.append(r"\subsection{" + line[5:] + "}")
        elif line.startswith(r"\# "):
            lines.append(r"\section{" + line[3:] + "}")
        else:
            lines.append(line)
    body = "\n".join(lines)
    return "\\documentclass{article}\n\\usepackage[utf8]{inputenc}\n\\begin{document}\n" + body + "\n\\end{document}\n"
