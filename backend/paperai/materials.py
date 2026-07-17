from __future__ import annotations

import csv
import hashlib
import io
import mimetypes
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree


TEXT_SUFFIXES = {".txt", ".md", ".rst", ".tex", ".py", ".r", ".json", ".yaml", ".yml", ".xml", ".html"}
TABLE_SUFFIXES = {".csv", ".tsv"}


def safe_filename(filename: str) -> str:
    name = Path(filename).name.strip().replace("\x00", "")
    name = re.sub(r"[^\w.()\- ]+", "_", name, flags=re.UNICODE)
    return name[:180] or "source"


def sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def extract_text(path: Path, media_type: str | None = None) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix in TABLE_SUFFIXES:
        delimiter = "\t" if suffix == ".tsv" else ","
        return _extract_delimited(path, delimiter)
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(paragraphs)
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        blocks: list[str] = []
        for sheet in workbook.worksheets:
            blocks.append(f"## Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                blocks.append(" | ".join("" if value is None else str(value) for value in row))
        return "\n".join(blocks)
    if suffix == ".pptx":
        return _extract_pptx(path)
    raise ValueError(f"Unsupported file type: {suffix or media_type or 'unknown'}")


def _extract_delimited(path: Path, delimiter: str) -> str:
    output = io.StringIO()
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        for index, row in enumerate(reader):
            if index >= 2000:
                output.write("\n[TRUNCATED: table exceeds 2000 rows]")
                break
            output.write(" | ".join(row) + "\n")
    return output.getvalue()


def _extract_pptx(path: Path) -> str:
    namespaces = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    slides: list[str] = []
    with zipfile.ZipFile(path) as archive:
        names = sorted(
            (name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
            key=lambda name: int(re.search(r"\d+", Path(name).stem).group()),
        )
        for index, name in enumerate(names, start=1):
            root = ElementTree.fromstring(archive.read(name))
            text = " ".join(node.text or "" for node in root.findall(".//a:t", namespaces))
            slides.append(f"## Slide {index}\n{text}")
    return "\n\n".join(slides)


def guessed_media_type(filename: str) -> str:
    return mimetypes.guess_type(filename)[0] or "application/octet-stream"
